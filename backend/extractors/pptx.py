"""PPTX extractor — preserves title + bullet structure per slide."""
import io
from pptx import Presentation


def extract_pptx(file_bytes: bytes) -> dict:
    try:
        prs = Presentation(io.BytesIO(file_bytes))
    except Exception as e:
        raise ValueError(f"Failed to open PPTX: {e}")

    pages = []
    for i, slide in enumerate(prs.slides, start=1):
        lines = []

        # ---- 1. Title (if the slide has a title placeholder with text) ----
        title_text = ""
        try:
            if slide.shapes.title and slide.shapes.title.text.strip():
                title_text = " ".join(slide.shapes.title.text.split())
        except Exception:
            title_text = ""

        if title_text:
            lines.append(f"📄 {title_text}")
            lines.append("")  # blank line after title

        # ---- 2. Body — collect per-paragraph, keep bullet structure ----
        body_lines = []
        for shape in slide.shapes:
            # Skip the title shape (already handled)
            try:
                if shape == slide.shapes.title:
                    continue
            except Exception:
                pass

            if not shape.has_text_frame:
                continue

            for para in shape.text_frame.paragraphs:
                para_text = "".join(run.text for run in para.runs).strip()
                if not para_text:
                    continue

                # Bullet detection: paragraph has bullet char OR level > 0
                is_bullet = False
                try:
                    # python-pptx doesn't expose bullet char directly,
                    # but pPr with buChar or buAutoNum implies a bullet
                    if para._pPr is not None:
                        if para._pPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}buChar') is not None or \
                           para._pPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}buAutoNum') is not None:
                            is_bullet = True
                except Exception:
                    pass

                # Also treat level > 0 as bullet
                if getattr(para, "level", 0) and para.level > 0:
                    is_bullet = True

                indent = "  " * (para.level or 0)

                if is_bullet:
                    body_lines.append(f"{indent}• {para_text}")
                else:
                    body_lines.append(f"{indent}{para_text}")

        # Add body if any
        if body_lines:
            lines.extend(body_lines)

        # Skip empty slides
        slide_text = "\n".join(lines).strip()
        if slide_text:
            pages.append({"page": i, "text": slide_text})

    full_text = "\n\n".join(p["text"] for p in pages if p["text"].strip())

    return {
        "text": full_text,
        "pages": pages,
        "source_type": "pptx",
        "used_ocr": False,
    }